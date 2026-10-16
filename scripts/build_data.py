"""Build the site's JSON data files from the seed workbook.

Run: python3 scripts/build_data.py path/to/jamaica-solar-companies-seed.xlsx path/to/bsj-approved-inverters.csv
Later this is replaced by reading from Supabase at build time.
"""
import sys, json, re, csv
import openpyxl

xlsx = sys.argv[1] if len(sys.argv) > 1 else "../jamaica-solar-companies-seed.xlsx"
bsj_csv = sys.argv[2] if len(sys.argv) > 2 else "../bsj-approved-inverters.csv"
wb = openpyxl.load_workbook(xlsx, data_only=True)

PARISHES = ['Kingston', 'St. Andrew', 'St. Thomas', 'Portland', 'St. Mary', 'St. Ann', 'Trelawny',
            'St. James', 'Hanover', 'Westmoreland', 'St. Elizabeth', 'Manchester', 'Clarendon', 'St. Catherine']

def slugify(s):
    s = s.lower().replace("&", "and")
    s = re.sub(r"\(.*?\)", "", s)
    return re.sub(r"[^a-z0-9]+", "-", s).strip("-")

def norm_parish(p):
    p = (p or "").strip()
    if not p: return None
    p = p.replace("St ", "St. ") if p.startswith("St ") else p
    return p if p in PARISHES else None

def tags_from(services, battery, btype):
    s = (services or "").lower()
    t = []
    if "pv" in s or "panel" in s or "solar" in s: t.append("Solar PV")
    if "batter" in s or "storage" in s or (battery or "").lower() == "yes" or "tesla" in s: t.append("Battery Storage")
    if "water heat" in s or "swh" in s: t.append("Water Heaters")
    if "mainten" in s: t.append("Maintenance")
    if "off-grid" in s or "off grid" in s: t.append("Off-grid")
    if "commercial" in s or "c&i" in s: t.append("Commercial")
    if "financ" in s or "hire purchase" in s: t.append("Financing")
    if "ev " in s or "ev charg" in s: t.append("EV Charging")
    return list(dict.fromkeys(t)) or ["Solar PV"]

def tagline(btype, services):
    b = (btype or "").lower()
    if "retail" in b: return "Solar equipment retailer"
    if "distributor" in b: return "Solar equipment distributor"
    if "seller" in b: return "Solar installation & equipment supply"
    if "developer" in b: return "Solar developer & installer"
    return "Solar installation"

def rows(ws):
    hdr = [c.value for c in ws[1]]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]: continue
        yield dict(zip(hdr, r))

providers = []
used = set()
for sheet, kind in (("Companies", "company"), ("Distributors & Retailers", "supplier")):
    for r in rows(wb[sheet]):
        name = r["company_name"].strip()
        slug = slugify(name)
        while slug in used: slug += "-2"
        used.add(slug)
        btype = r.get("business_type") or ""
        parish = norm_parish(r.get("parish_hq"))
        served = [p for p in (r.get("parishes_served") or "").replace(",", ";").split(";") if norm_parish(p)]
        served = [norm_parish(p) for p in served]
        parish_list = served or ([parish] if parish else [])
        sources = [s.strip() for s in (r.get("sources") or "").split(";") if s.strip()]
        k = kind
        if kind == "company" and ("seller" in btype.lower()):
            k = "company"  # installs and sells
        providers.append({
            "slug": slug,
            "kind": k,
            "name": name,
            "businessType": btype,
            "tagline": tagline(btype, r.get("services")),
            "tags": tags_from(r.get("services"), r.get("battery_storage"), btype),
            "servicesRaw": r.get("services") or "",
            "city": r.get("city_area") or "",
            "parishHq": parish,
            "parishes": parish_list,
            "address": r.get("address") or "",
            "website": r.get("website") or "",
            "phone": r.get("phone") or "",
            "email": r.get("email") or "",
            "batteryStorage": (r.get("battery_storage") or "").lower() == "yes",
            "sponsored": False,
            "verified": (r.get("status") or "").lower().startswith("verified"),
            "checked": r.get("date_verified") or "",
            "sources": sources,
            "notes": r.get("notes") or "",
        })

providers.sort(key=lambda p: p["name"].lower())

lenders = []
ws = wb["Homeowner Lenders"]
hdr = [c.value for c in ws[2]]
for r in ws.iter_rows(min_row=3, values_only=True):
    if not r[0] or str(r[0]).startswith(("Incentive", "Solar PV income", "Example", "Loan amount", "Annual rate")):
        if r[0] and str(r[0]).startswith("Incentive"): break
        continue
    if isinstance(r[0], (int, float)): continue
    d = dict(zip(hdr, r))
    lenders.append({k: ("" if v is None else v) for k, v in d.items()})

inverters = []
with open(bsj_csv, encoding="utf-8") as f:
    for r in csv.DictReader(f):
        inverters.append({
            "brand": r["brand"], "model": r["model"] or "(model name missing in BSJ list)", "key": r["lookup_key"],
            "kw": float(r["output_kw"]) if r["output_kw"] else None, "output": r["output_raw"],
            "voltage": r["ac_output_v"], "size": r["size_class"], "hybrid": bool(r["hybrid_battery_capable"]),
            "report": r["bsj_report_no"], "status": r["approval_status"], "date": r["approval_date"],
            "issue": r["data_issue"],
        })

json.dump(providers, open("src/data/providers.json", "w"), indent=1, ensure_ascii=False)
json.dump(lenders, open("src/data/lenders.json", "w"), indent=1, ensure_ascii=False)
json.dump(inverters, open("src/data/inverters.json", "w"), ensure_ascii=False)
print(len(providers), "providers;", sum(p["kind"] == "supplier" for p in providers), "suppliers;",
      len(lenders), "lenders;", len(inverters), "inverters")
